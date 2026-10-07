"""Document Ingestion & Chunking Pipeline.

Extracts text from PDFs, PPTXs, and text notes.
Preserves metadata: source filename, page/slide number, section, document type.
Performs clean-up, token-approximated chunking with configurable overlap.
"""

import os
import re
import json
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
from pydantic import BaseModel, Field


class DocumentChunk(BaseModel):
    """Schema for a chunk of syllabus / reference text."""
    chunk_id: str = Field(..., description="Unique deterministic chunk ID")
    content: str = Field(..., description="Cleaned chunk text content")
    source: str = Field(..., description="Original filename, e.g. OS_Deadlocks.pptx")
    page: Optional[int] = Field(None, description="Page or slide number (1-indexed)")
    doc_type: str = Field("unknown", description="pdf, pptx, txt, etc.")
    topic: Optional[str] = Field(None, description="Topic tag or heading")
    citation: str = Field(..., description="Human/LLM-readable citation string")
    metadata: Dict[str, Any] = Field(default_factory=dict)


class IngestionReport(BaseModel):
    """Summary report after running document ingestion."""
    total_documents: int
    total_chunks: int
    files_processed: List[str]
    skipped_files: List[str]
    output_chunks_file: str
    output_catalog_file: str


def clean_text(raw_text: str) -> str:
    """Cleans extracted text: normalizes whitespace, fixes line-break hyphens."""
    if not raw_text:
        return ""
    # Fix hyphenated line breaks (e.g. "dead-\nlock" -> "deadlock")
    text = re.sub(r"(\w+)-\s*\n\s*(\w+)", r"\1\2", raw_text)
    # Replace non-breaking spaces and excessive whitespace
    text = text.replace("\xa0", " ")
    # Remove isolated page number headers/footers e.g. "Page 12 of 45" or "--- 12 ---"
    text = re.sub(r"(?i)\bpage\s+\d+\s+of\s+\d+\b", "", text)
    text = re.sub(r"-+\s*\d+\s*-+", "", text)
    # Normalize multiple whitespace and newlines
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
    clean_lines = [line for line in lines if line]
    return "\n".join(clean_lines)


def format_citation(source: str, page: Optional[int], doc_type: str) -> str:
    """Formats student-friendly citation string."""
    if page is None:
        return f"[Source: {source}]"
    unit = "slide" if doc_type.lower() in ("pptx", "ppt") else "page"
    return f"[Source: {source}, {unit} {page}]"


def extract_text_from_pdf(file_path: Path) -> List[Tuple[int, str]]:
    """Extracts text per page from a PDF file using pypdf.

    Returns: list of (page_number, page_text)
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ImportError("pypdf is required to parse PDF files. Install via `pip install pypdf`.")

    reader = PdfReader(str(file_path))
    pages_text = []
    for idx, page in enumerate(reader.pages):
        page_num = idx + 1
        text = page.extract_text() or ""
        cleaned = clean_text(text)
        if cleaned:
            pages_text.append((page_num, cleaned))
    return pages_text


def extract_text_from_pptx(file_path: Path) -> List[Tuple[int, str]]:
    """Extracts text per slide from a PowerPoint (.pptx) file using python-pptx.

    Returns: list of (slide_number, slide_text)
    """
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError("python-pptx is required to parse PPTX files. Install via `pip install python-pptx`.")

    prs = Presentation(str(file_path))
    slides_text = []
    for idx, slide in enumerate(prs.slides):
        slide_num = idx + 1
        text_runs = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    p_text = paragraph.text.strip()
                    if p_text:
                        text_runs.append(p_text)
        slide_content = clean_text("\n".join(text_runs))
        if slide_content:
            slides_text.append((slide_num, slide_content))
    return slides_text


def extract_text_from_txt(file_path: Path) -> List[Tuple[int, str]]:
    """Extracts text from a plain text file."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()
    cleaned = clean_text(content)
    return [(1, cleaned)] if cleaned else []


def extract_document(file_path: Path) -> List[Tuple[int, str, str]]:
    """Dispatches extraction by file suffix.

    Returns: list of (page_num, page_text, doc_type)
    """
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return [(p, t, "pdf") for p, t in extract_text_from_pdf(file_path)]
    elif suffix in (".pptx", ".ppt"):
        return [(p, t, "pptx") for p, t in extract_text_from_pptx(file_path)]
    elif suffix in (".txt", ".md"):
        return [(p, t, "txt") for p, t in extract_text_from_txt(file_path)]
    else:
        raise ValueError(f"Unsupported document format: {suffix}")


def chunk_document(
    pages: List[Tuple[int, str, str]],
    source_filename: str,
    chunk_size: int = 400,
    overlap: int = 50,
    topic: Optional[str] = None
) -> List[DocumentChunk]:
    """Splits document pages/slides into overlapping DocumentChunk objects."""
    chunks: List[DocumentChunk] = []
    step = max(1, chunk_size - overlap)

    for page_num, text, doc_type in pages:
        words = text.split()
        if not words:
            continue

        if len(words) <= chunk_size:
            chunk_content = " ".join(words)
            chunk_id = f"{source_filename}_p{page_num}_c0"
            citation = format_citation(source_filename, page_num, doc_type)
            chunks.append(
                DocumentChunk(
                    chunk_id=chunk_id,
                    content=chunk_content,
                    source=source_filename,
                    page=page_num,
                    doc_type=doc_type,
                    topic=topic,
                    citation=citation,
                    metadata={"word_count": len(words), "page": page_num}
                )
            )
        else:
            i = 0
            chunk_idx = 0
            while i < len(words):
                chunk_words = words[i:i + chunk_size]
                chunk_content = " ".join(chunk_words)
                chunk_id = f"{source_filename}_p{page_num}_c{chunk_idx}"
                citation = format_citation(source_filename, page_num, doc_type)
                chunks.append(
                    DocumentChunk(
                        chunk_id=chunk_id,
                        content=chunk_content,
                        source=source_filename,
                        page=page_num,
                        doc_type=doc_type,
                        topic=topic,
                        citation=citation,
                        metadata={"word_count": len(chunk_words), "page": page_num, "sub_chunk": chunk_idx}
                    )
                )
                chunk_idx += 1
                i += step

    return chunks


def ingest_raw_documents(
    raw_dir: Path,
    processed_dir: Path,
    metadata_dir: Path,
    chunk_size: int = 400,
    overlap: int = 50
) -> Tuple[List[DocumentChunk], IngestionReport]:
    """Scans raw_dir for PDFs/PPTXs, extracts, chunks, and saves to processed_dir and metadata_dir."""
    raw_dir = Path(raw_dir)
    processed_dir = Path(processed_dir)
    metadata_dir = Path(metadata_dir)

    raw_dir.mkdir(parents=True, exist_ok=True)
    processed_dir.mkdir(parents=True, exist_ok=True)
    metadata_dir.mkdir(parents=True, exist_ok=True)

    all_chunks: List[DocumentChunk] = []
    processed_files: List[str] = []
    skipped_files: List[str] = []
    catalog: Dict[str, Any] = {}

    for file_path in sorted(raw_dir.iterdir()):
        if file_path.is_file() and not file_path.name.startswith("."):
            suffix = file_path.suffix.lower()
            if suffix in (".pdf", ".pptx", ".ppt", ".txt", ".md"):
                try:
                    pages = extract_document(file_path)
                    doc_chunks = chunk_document(
                        pages=pages,
                        source_filename=file_path.name,
                        chunk_size=chunk_size,
                        overlap=overlap,
                        topic=file_path.stem.replace("_", " ").title()
                    )
                    all_chunks.extend(doc_chunks)
                    processed_files.append(file_path.name)
                    catalog[file_path.name] = {
                        "path": str(file_path),
                        "suffix": suffix,
                        "pages_extracted": len(pages),
                        "chunks_created": len(doc_chunks)
                    }
                except Exception as e:
                    skipped_files.append(f"{file_path.name} (Error: {e})")
            else:
                skipped_files.append(f"{file_path.name} (Unsupported format)")

    # Save chunks to processed_dir / chunks.json
    chunks_path = processed_dir / "chunks.json"
    with open(chunks_path, "w", encoding="utf-8") as f:
        json.dump([c.model_dump() for c in all_chunks], f, indent=2)

    # Save catalog to metadata_dir / catalog.json
    catalog_path = metadata_dir / "catalog.json"
    with open(catalog_path, "w", encoding="utf-8") as f:
        json.dump({
            "total_documents": len(processed_files),
            "total_chunks": len(all_chunks),
            "documents": catalog
        }, f, indent=2)

    report = IngestionReport(
        total_documents=len(processed_files),
        total_chunks=len(all_chunks),
        files_processed=processed_files,
        skipped_files=skipped_files,
        output_chunks_file=str(chunks_path),
        output_catalog_file=str(catalog_path)
    )
    return all_chunks, report
